"""Build a text reader PDF directly from Office XML, entirely in Python.

DOCX pages are generated reader pages, not a reconstruction of Word's layout.
PPTX keeps one reader page per slide, including blank/hidden slides, without notes.
Evidence is subsequently extracted from these exact PDF bytes.
"""
from pathlib import Path
import threading
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table as WordTable
from pptx import Presentation
from reportlab import rl_config
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.models import FileFormat

_font_lock = threading.Lock()


def reader_font():
    name = 'DocuMindOfficeReader'
    with _font_lock:
        if name not in pdfmetrics.getRegisteredFontNames():
            candidates = [Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
                          Path('C:/Windows/Fonts/arial.ttf'),
                          Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
            candidates.extend(Path(folder) / 'Vera.ttf' for folder in rl_config.TTFSearchPath)
            for path in candidates:
                if path.is_file():
                    pdfmetrics.registerFont(TTFont(name, str(path)))
                    break
            else:
                raise ValueError('No Unicode reader font is available.')
    return name


def _paragraph(text, style):
    return Paragraph(escape(text).replace('\n', '<br/>').replace('\t', '    '), style)


def _paragraph_parts(paragraph):
    """Keep explicit page breaks and hyperlink text in original XML order."""
    pieces = []
    for element in paragraph._p.iter():
        if element.tag == qn('w:t'):
            pieces.append(element.text or '')
        elif element.tag == qn('w:tab'):
            pieces.append('\t')
        elif element.tag in (qn('w:br'), qn('w:cr')):
            if element.get(qn('w:type')) == 'page':
                yield ''.join(pieces)
                yield None
                pieces = []
            else:
                pieces.append('\n')
    yield ''.join(pieces)


def _render_docx(source, output):
    document = WordDocument(source)
    section = document.sections[0]
    width = section.page_width.pt if section.page_width else 595.28
    height = section.page_height.pt if section.page_height else 841.89
    left = min(section.left_margin.pt if section.left_margin else 54, width / 4)
    right = min(section.right_margin.pt if section.right_margin else 54, width / 4)
    top = min(section.top_margin.pt if section.top_margin else 54, height / 4)
    bottom = min(section.bottom_margin.pt if section.bottom_margin else 54, height / 4)
    font = reader_font()
    style = ParagraphStyle('OfficeBody', fontName=font, fontSize=11, leading=15,
                           spaceAfter=6, splitLongWords=True)
    table_style = ParagraphStyle('OfficeCell', parent=style, fontSize=9, leading=12, spaceAfter=2)
    story = []
    for block in document.iter_inner_content():
        if isinstance(block, WordTable):
            rows, seen = [], set()
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    # Merged cells can appear repeatedly in the row/cell list.
                    text = '' if cell._tc in seen else cell.text
                    seen.add(cell._tc)
                    cells.append(_paragraph(text, table_style))
                rows.append(cells)
            if rows and rows[0]:
                columns = max(map(len, rows))
                rows = [row + [''] * (columns - len(row)) for row in rows]
                story.append(LongTable(rows, colWidths=[(width-left-right-12)/columns]*columns,
                                       splitByRow=1, splitInRow=1, style=[
                                           ('VALIGN', (0,0), (-1,-1), 'TOP'),
                                           ('GRID', (0,0), (-1,-1), .4, '#cccccc')]))
                story.append(Spacer(1,6))
            continue
        if block.paragraph_format.page_break_before:
            story.append(PageBreak())
        for text in _paragraph_parts(block):
            if text is None:
                story.append(PageBreak())
            elif text:
                story.append(_paragraph(text, style))
    if not story:
        story.append(Spacer(1,1))
    SimpleDocTemplate(str(output), pagesize=(width,height), leftMargin=left,
                      rightMargin=right, topMargin=top, bottomMargin=bottom).build(story)


def _shape_texts(shapes):
    for shape in shapes:
        if hasattr(shape, 'shapes'):
            yield from _shape_texts(shape.shapes)
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text:
                    yield paragraph.text
        if shape.has_table:
            for row in shape.table.rows:
                yield '    '.join(cell.text for cell in row.cells if not cell.is_spanned)
        if getattr(shape, 'has_chart', False) and shape.chart.has_title:
            title = shape.chart.chart_title
            if title.has_text_frame and title.text_frame.text:
                yield title.text_frame.text


def _render_pptx(source, output):
    presentation = Presentation(source)
    width, height = presentation.slide_width.pt, presentation.slide_height.pt
    font = reader_font()
    margin = min(30, width/10, height/10)
    available_width, available_height = width-2*margin, height-2*margin
    canvas = Canvas(str(output), pagesize=(width,height))
    for slide in presentation.slides:
        texts = list(_shape_texts(slide.shapes))
        size = 14.0
        for _ in range(30):
            style = ParagraphStyle('SlideReader', fontName=font, fontSize=size,
                                   leading=size*1.3, spaceAfter=size*.35, splitLongWords=True)
            paragraphs = [_paragraph(text, style) for text in texts]
            heights = [p.wrap(available_width, 1000000)[1] for p in paragraphs]
            total = sum(heights) + max(0,len(heights)-1)*style.spaceAfter
            if total <= available_height:
                break
            size *= min(.9, available_height/total*.95)
        else:
            raise ValueError('The slide text cannot fit in its reader page.')
        y = height-margin
        for paragraph, h in zip(paragraphs,heights):
            y -= h
            paragraph.drawOn(canvas,margin,y)
            y -= style.spaceAfter
        canvas.showPage()  # Even blank/hidden slides retain their number.
    canvas.save()


def render_office_pdf(source, file_format, output_dir):
    output = Path(output_dir) / 'office-reader.pdf'
    if file_format == FileFormat.DOCX:
        _render_docx(source, output)
    elif file_format == FileFormat.PPTX:
        _render_pptx(source, output)
    else:
        raise ValueError('Unsupported Office format.')
    return output
