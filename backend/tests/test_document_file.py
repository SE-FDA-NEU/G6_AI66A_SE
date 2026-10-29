"""Tests for GET /documents/{document_id}/file and rendered-file responses.

Verifies:
- Valid files (PDF, DOCX, PPTX) return the correct readable representation.
- Page metadata headers (X-Page-Count, X-Page-Kind, X-Initial-Page, etc.) are present.
- Storage paths are never leaked in headers or error bodies.
- Ownership protection from US05: 403 'Access denied' on cross-account requests.
- Missing documents and missing physical files return 404.
- Page outside document bounds returns 422.
- Unauthenticated requests return 401.
"""
from datetime import datetime, timedelta
from pathlib import Path
from io import BytesIO

import pytest
from docx import Document as WordDocument
from pptx import Presentation
from pypdf import PdfReader

from app.core.security import create_access_token
from app.models import Document, FileFormat
from app.services.reader_extraction import extract_document_for_reader
from test_pdf_extraction import make_pdf


@pytest.mark.parametrize('endpoint', ['file', 'rendered-file'])
@pytest.mark.parametrize('page', [
    pytest.param(1, id='TC-US09-05-first-citation-page'),
    pytest.param(3, id='TC-US09-06-last-citation-page'),
])
def test_citation_page_selects_readable_original_without_changing_source(
        client, session_factory, test_files, endpoint, page):
    source = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, source, FileFormat.PDF)
    response = client.get(f'/documents/{doc_id}/{endpoint}', params={'page': page})
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/pdf'
    assert response.headers['x-initial-page'] == str(page)
    assert response.headers['x-page-count'] == '3'
    assert response.content == source.read_bytes()
    assert PdfReader(BytesIO(response.content)).pages[page - 1].extract_text().strip() == f'Page {page} text'


@pytest.mark.parametrize('endpoint', ['file', 'rendered-file'])
@pytest.mark.parametrize('page,detail', [
    pytest.param('0', 'Page must be greater than or equal to 1', id='TC-US09-07-page-zero'),
    pytest.param('-1', 'Page must be greater than or equal to 1', id='TC-US09-08-negative-page'),
    pytest.param('4', 'Page outside document', id='TC-US09-09-page-past-last'),
    pytest.param('invalid', None, id='TC-US09-10-nonnumeric-page'),
    pytest.param('1.5', None, id='TC-US09-11-fractional-page'),
])
def test_invalid_citation_page_is_rejected_without_serving_file(
        client, session_factory, test_files, endpoint, page, detail):
    source = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, source, FileFormat.PDF)
    response = client.get(f'/documents/{doc_id}/{endpoint}', params={'page': page})
    assert response.status_code == 422
    assert response.headers['content-type'] == 'application/json'
    assert str(source) not in response.text
    assert 'x-initial-page' not in response.headers
    if detail:
        assert response.json()['detail'] == detail
    else:
        assert any(error['loc'] == ['query', 'page'] for error in response.json()['detail'])


@pytest.mark.parametrize('content', [
    pytest.param(b'q 100 0 0 100 72 600 cm /Im0 Do Q', id='TC-US09-12-image-only-page-readable'),
    pytest.param(b'', id='TC-US09-13-blank-page-readable'),
])
def test_textless_page_remains_readable_but_cannot_accept_citation_span(
        client, session_factory, tmp_path, content):
    source = make_pdf(tmp_path / 'textless.pdf', [
        b'BT /F1 12 Tf 72 720 Td (Text on first page) Tj ET', content,
    ])
    doc_id = create_document(session_factory, 123, source, FileFormat.PDF, page_count=2)
    extract_document_for_reader(doc_id, session_factory)
    response = client.get(f'/documents/{doc_id}/pages')
    assert response.status_code == 200
    metadata = response.json()
    assert metadata['pages'][1]['page'] == 2
    assert metadata['pages'][1]['text'] == ''
    assert metadata['pages'][1]['positions'] == []
    for endpoint in ('file', 'rendered-file'):
        readable = client.get(f'/documents/{doc_id}/{endpoint}', params={'page': 2})
        assert readable.status_code == 200
        assert readable.headers['x-initial-page'] == '2'
        assert readable.content == source.read_bytes()
        pdf = PdfReader(BytesIO(readable.content))
        assert len(pdf.pages) == 2 and pdf.pages[1].extract_text() == ''
        if content:
            assert len(pdf.pages[1].images) == 1
    # Textless pages are valid source pages; invented text spans are not.
    invalid = client.get(f'/documents/{doc_id}/reader.pdf', params={
        'version': metadata['reader_version'], 'highlight_page': 2, 'start': 0, 'end': 1,
    })
    assert invalid.status_code == 422
    assert invalid.json()['detail'] == 'Highlight span is outside the page text.'


@pytest.fixture
def test_files(tmp_path):
    # 1. PDF file
    pdf_path = make_pdf(tmp_path / 'sample.pdf', [
        b'BT /F1 12 Tf 72 720 Td (Page 1 text) Tj ET',
        b'BT /F1 12 Tf 72 720 Td (Page 2 text) Tj ET',
        b'BT /F1 12 Tf 72 720 Td (Page 3 text) Tj ET',
    ])

    # 2. DOCX file
    docx_path = tmp_path / 'sample.docx'
    doc = WordDocument()
    doc.add_paragraph('Hello world from docx')
    doc.save(docx_path)

    # 3. PPTX file
    pptx_path = tmp_path / 'sample.pptx'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(pptx_path)

    return {
        FileFormat.PDF: pdf_path,
        FileFormat.DOCX: docx_path,
        FileFormat.PPTX: pptx_path,
    }


def create_document(session_factory, owner_id, file_path, file_format, page_count=3, filename=None):
    with session_factory() as db:
        doc = Document(
            owner_id=owner_id,
            original_filename=filename or file_path.name,
            storage_path=str(file_path),
            file_format=file_format,
            file_size_bytes=file_path.stat().st_size if file_path.exists() else 100,
            page_count=page_count,
            extracted_text=[{'page': i, 'text': f'Text page {i}'} for i in range(1, page_count + 1)],
            last_accessed_at=datetime(2026, 1, 1),
            expires_at=datetime(2026, 1, 1) + timedelta(days=30),
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)
        return doc.id


def test_unauthenticated_request_returns_401(client, session_factory, test_files):
    doc_id = create_document(session_factory, 123, test_files[FileFormat.PDF], FileFormat.PDF)
    res = client.get(f'/documents/{doc_id}/file', headers={'Authorization': ''})
    assert res.status_code == 401


def test_valid_pdf_returns_file_and_metadata_headers(client, session_factory, test_files):
    pdf_path = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, pdf_path, FileFormat.PDF, page_count=3, filename='my_contract.pdf')

    res = client.get(f'/documents/{doc_id}/file')
    assert res.status_code == 200
    assert res.headers['content-type'] == 'application/pdf'
    assert 'inline' in res.headers.get('content-disposition', '')
    assert 'my_contract.pdf' in res.headers.get('content-disposition', '')

    # Verify page metadata headers
    assert res.headers['x-page-count'] == '3'
    assert res.headers['x-page-kind'] == 'page'
    assert res.headers['x-original-filename'] == 'my_contract.pdf'
    assert res.headers['x-file-format'] == 'PDF'

    # Verify content matches original bytes
    assert res.content == pdf_path.read_bytes()

    # Verify storage path is NEVER exposed in headers or body
    for header_val in res.headers.values():
        assert str(pdf_path) not in header_val
    assert str(pdf_path).encode() not in res.content

    # Verify retention update (BR06)
    with session_factory() as db:
        updated = db.get(Document, doc_id)
        assert updated.last_accessed_at > datetime(2026, 1, 1)
        assert (updated.expires_at - updated.last_accessed_at).days == 30


def test_page_selection_and_validation(client, session_factory, test_files):
    pdf_path = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, pdf_path, FileFormat.PDF, page_count=3)

    # Valid page
    res = client.get(f'/documents/{doc_id}/file?page=2')
    assert res.status_code == 200
    assert res.headers['x-initial-page'] == '2'

    # Page outside document bounds returns 422
    res_exceeded = client.get(f'/documents/{doc_id}/file?page=10')
    assert res_exceeded.status_code == 422
    assert res_exceeded.json()['detail'] == 'Page outside document'

    # Zero or negative page returns 422
    assert client.get(f'/documents/{doc_id}/file?page=0').status_code == 422
    assert client.get(f'/documents/{doc_id}/file?page=-1').status_code == 422


def test_valid_docx_and_pptx_files(client, session_factory, test_files):
    # DOCX
    docx_path = test_files[FileFormat.DOCX]
    docx_id = create_document(session_factory, 123, docx_path, FileFormat.DOCX, page_count=2, filename='notes.docx')
    res_docx = client.get(f'/documents/{docx_id}/file')
    assert res_docx.status_code == 200
    assert 'wordprocessingml.document' in res_docx.headers['content-type']
    assert res_docx.headers['x-page-kind'] == 'page'
    assert res_docx.headers['x-file-format'] == 'DOCX'
    assert res_docx.content == docx_path.read_bytes()

    # PPTX
    pptx_path = test_files[FileFormat.PPTX]
    pptx_id = create_document(session_factory, 123, pptx_path, FileFormat.PPTX, page_count=5, filename='deck.pptx')
    res_pptx = client.get(f'/documents/{pptx_id}/file')
    assert res_pptx.status_code == 200
    assert 'presentationml.presentation' in res_pptx.headers['content-type']
    assert res_pptx.headers['x-page-kind'] == 'slide'
    assert res_pptx.headers['x-file-format'] == 'PPTX'
    assert res_pptx.content == pptx_path.read_bytes()


def test_rendered_file_response(client, session_factory, test_files, tmp_path):
    # For PDF, rendered file is the PDF itself
    pdf_path = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, pdf_path, FileFormat.PDF, page_count=3, filename='test.pdf')

    res_rendered = client.get(f'/documents/{doc_id}/file?rendered=true')
    assert res_rendered.status_code == 200
    assert res_rendered.headers['content-type'] == 'application/pdf'
    assert res_rendered.content == pdf_path.read_bytes()

    # Alias /rendered-file endpoint
    res_alias = client.get(f'/documents/{doc_id}/rendered-file')
    assert res_alias.status_code == 200
    assert res_alias.headers['content-type'] == 'application/pdf'
    assert res_alias.content == pdf_path.read_bytes()


def test_missing_document_returns_404(client):
    res = client.get('/documents/999999/file')
    assert res.status_code == 404
    assert res.json() == {'detail': 'Document not found'}


def test_missing_physical_file_returns_404_without_path_leak(client, session_factory, tmp_path):
    ghost_path = tmp_path / 'non_existent_file.pdf'
    doc_id = create_document(session_factory, 123, ghost_path, FileFormat.PDF, filename='ghost.pdf')

    res = client.get(f'/documents/{doc_id}/file')
    assert res.status_code == 404
    assert res.json() == {'detail': 'File not found'}

    # Ensure storage path is NOT leaked
    assert str(ghost_path) not in res.text
    for val in res.headers.values():
        assert str(ghost_path) not in val


def test_cross_account_request_denied_with_403(client, session_factory, test_files):
    # Document owned by User 123
    pdf_path = test_files[FileFormat.PDF]
    doc_id = create_document(session_factory, 123, pdf_path, FileFormat.PDF, filename='secret.pdf')

    # User 124 tries to access User 123's document
    other_token = f"Bearer {create_access_token({'sub': '124'})}"
    res = client.get(f'/documents/{doc_id}/file', headers={'Authorization': other_token})

    # Exact generic error from US05
    assert res.status_code == 403
    assert res.json() == {'detail': 'Access denied'}

    # Ensure no content or path is leaked
    assert str(pdf_path) not in res.text
    assert 'secret.pdf' not in res.text

    # Verify retention timestamps for User 123's document remain untouched
    with session_factory() as db:
        doc = db.get(Document, doc_id)
        assert doc.last_accessed_at == datetime(2026, 1, 1)
        assert doc.expires_at == datetime(2026, 1, 1) + timedelta(days=30)
