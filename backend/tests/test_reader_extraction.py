import hashlib
from datetime import datetime, timedelta

import pytest
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy.exc import SQLAlchemyError

from app.core.security import create_access_token
from app.models import Document, FileFormat
from app.services import reader_extraction as reader
from app.services.pdf_extraction import extract_pdf
from test_pdf_extraction import make_pdf


@pytest.fixture
def sources(tmp_path):
    word = WordDocument()
    word.add_paragraph('DOCX page one alpha')
    word.add_paragraph('DOCX page one beta')
    word.add_page_break()
    word.add_paragraph('DOCX page two gamma')
    word.add_paragraph('DOCX page two delta')
    docx = tmp_path / 'source.docx'
    word.save(docx)
    slides = Presentation()
    for i in range(3):
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        if i < 2:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = f'PPTX slide {i + 1}'
        if i == 1:
            slide._element.set('show', '0')
        slide.notes_slide.notes_text_frame.text = 'SPEAKER NOTES MUST NOT BECOME READER PAGES'
    pptx = tmp_path / 'source.pptx'
    slides.save(pptx)
    return {FileFormat.DOCX: docx, FileFormat.PPTX: pptx}


def save_document(factory, source, fmt):
    with factory() as db:
        document = Document(owner_id=123, original_filename=source.name, storage_path=str(source),
                            file_format=fmt, file_size_bytes=source.stat().st_size,
                            expires_at=datetime.now() + timedelta(days=30))
        db.add(document)
        db.commit()
        return document.id


@pytest.fixture
def keyword_reader(session_factory, tmp_path):
    """Fixed physical pages and offsets shared by extraction/presentation tests."""
    texts = ['Alpha first page.', 'Middle keyword keyword.', 'Last page Omega']
    source = make_pdf(tmp_path / 'keyword-reader.pdf', [
        f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode('ascii') for text in texts
    ])
    document_id = save_document(session_factory, source, FileFormat.PDF)
    result = reader.extract_document_for_reader(document_id, session_factory)
    return document_id, source, result['extracted_text']


@pytest.mark.parametrize('query,page_number,spans', [
    pytest.param('keyword', 2, [(7, 14), (15, 22)], id='TC-US04-19-known-page-keyword'),
    pytest.param('Alpha', 1, [(0, 5)], id='TC-US04-20-first-page-first-character'),
    pytest.param('Omega', 3, [(10, 15)], id='TC-US04-21-last-page-last-character'),
])
def test_reader_keyword_search_returns_known_page_and_source_positions(
        client, keyword_reader, query, page_number, spans):
    document_id, _, pages = keyword_reader
    response = client.get('/search', params={'q': query})
    assert response.status_code == 200
    result = response.json()
    assert result['query'] == query
    assert result['total_hits'] == len(result['hits']) == 1
    hit = result['hits'][0]
    assert hit['document_id'] == document_id
    assert hit['page'] == page_number
    page_response = client.get(f'/documents/{document_id}/pages/{page_number}/text')
    assert page_response.status_code == 200
    page = page_response.json()
    assert page['page'] == page_number
    assert page['text'] == pages[page_number - 1]['text']
    # Search hit offsets identify the source passage, not just the keyword.
    assert (hit['start'], hit['end']) == (0, len(page['text']))
    assert page['text'][hit['start']:hit['end']] == hit['snippet']
    assert hit['reader_version'] == pages[page_number - 1]['reader']['sha256']
    assert hit['highlight_url'] == (
        f'/documents/{document_id}?page={page_number}&start=0&end={len(page["text"])}')
    for start, end in spans:
        assert page['text'][start:end] == query
        positions = [position for position in page['positions']
                     if position['start'] < end and position['end'] > start]
        assert positions
        assert positions[0]['start'] == start
        assert positions[-1]['end'] == end
        assert ''.join(position['text'] for position in positions) == query


@pytest.mark.parametrize('query', [pytest.param('absentword', id='TC-US04-22-no-matches')])
def test_reader_search_no_match_is_a_valid_empty_result(client, keyword_reader, query):
    response = client.get('/search', params={'q': query})
    assert response.status_code == 200
    assert response.json() == {'query': query, 'total_hits': 0, 'hits': []}


@pytest.mark.parametrize('params,error_location', [
    pytest.param({}, ['query', 'q'], id='TC-US04-23-missing-query'),
    pytest.param({'q': ''}, ['query', 'q'], id='TC-US04-24-empty-query'),
    pytest.param({'q': 'x' * 501}, ['query', 'q'], id='TC-US04-25-overlong-query'),
    pytest.param({'q': ' \t\n'}, None, id='TC-US04-26-whitespace-query'),
    pytest.param({'q': '!!!'}, None, id='TC-US04-27-punctuation-query'),
])
def test_reader_search_invalid_query_is_not_a_no_match_result(client, keyword_reader, params, error_location):
    response = client.get('/search', params=params)
    assert response.status_code == 422
    payload = response.json()
    assert 'hits' not in payload and 'total_hits' not in payload
    if error_location:
        assert any(error['loc'] == error_location for error in payload['detail'])
    else:
        assert payload['detail'] == 'Search query must contain a word or number.'


@pytest.mark.parametrize('page_number,expected_text', [
    pytest.param(1, 'Alpha first page.', id='TC-US04-28-first-page-text'),
    pytest.param(2, 'Middle keyword keyword.', id='TC-US04-29-middle-page-text'),
    pytest.param(3, 'Last page Omega', id='TC-US04-30-last-page-text'),
])
def test_reader_page_text_preserves_physical_page_and_positions(client, keyword_reader, page_number, expected_text):
    document_id, _, pages = keyword_reader
    response = client.get(f'/documents/{document_id}/pages/{page_number}/text')
    assert response.status_code == 200
    payload = response.json()
    assert payload == {'document_id': document_id, 'page': page_number,
                       'text': expected_text, 'positions': pages[page_number - 1]['positions']}
    assert payload['positions']
    assert all(expected_text[position['start']:position['end']] == position['text']
               for position in payload['positions'])


@pytest.mark.parametrize('page_input,status,detail', [
    pytest.param('0', 404, 'Page not found', id='TC-US04-31-page-zero'),
    pytest.param('-1', 404, 'Page not found', id='TC-US04-32-negative-page'),
    pytest.param('4', 404, 'Page not found', id='TC-US04-33-one-past-last-page'),
    pytest.param('invalid', 422, None, id='TC-US04-34-nonnumeric-page'),
    pytest.param('1.5', 422, None, id='TC-US04-35-fractional-page'),
])
def test_reader_page_text_rejects_invalid_or_missing_pages(client, keyword_reader, page_input, status, detail):
    document_id, _, _ = keyword_reader
    response = client.get(f'/documents/{document_id}/pages/{page_input}/text')
    assert response.status_code == status
    payload = response.json()
    assert 'text' not in payload and 'positions' not in payload
    if detail:
        assert payload['detail'] == detail
    else:
        assert any(error['loc'] == ['path', 'page'] for error in payload['detail'])


def fake_converter(source, fmt, output_dir):
    # Test publication/rollback independently of the native Office renderer.
    texts = ['DOCX page one alpha DOCX page one beta', 'DOCX page two gamma DOCX page two delta'] if fmt == FileFormat.DOCX else ['PPTX slide 1', 'PPTX slide 2', '']
    return make_pdf(output_dir / 'source.pdf', [
        f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode() if text else b'' for text in texts
    ])


def assert_reader_matches_api(client, document_id, result, tmp_path):
    response = client.get(f'/documents/{document_id}/pages')
    assert response.status_code == 200
    metadata = response.json()
    file_response = client.get(metadata['reader_url'])
    assert file_response.status_code == 200
    assert file_response.headers['content-type'] == 'application/pdf'
    assert file_response.headers['x-reader-version'] == metadata['reader_version']
    assert hashlib.sha256(file_response.content).hexdigest() == metadata['reader_version']
    displayed = tmp_path / 'actually-displayed.pdf'
    displayed.write_bytes(file_response.content)
    displayed_pages = extract_pdf(displayed)['extracted_text']
    assert metadata['page_count'] == result['page_count'] == len(displayed_pages)
    for extracted, shown in zip(metadata['pages'], displayed_pages):
        assert extracted['page'] == shown['page']
        assert extracted['text'] == shown['text']
        assert extracted['positions'] == shown['positions']
        assert extracted['reader']['sha256'] == metadata['reader_version']
    return metadata, file_response.content


@pytest.mark.parametrize('fmt', [FileFormat.DOCX, FileFormat.PPTX])
def test_persisted_pages_are_extracted_from_the_exact_reader_file(client, session_factory, sources, fmt, tmp_path):
    document_id = save_document(session_factory, sources[fmt], fmt)
    calls = []
    def convert(*args):
        calls.append(1)
        return fake_converter(*args)
    result = reader.extract_document_for_reader(document_id, session_factory, convert)
    expected_pages = 2 if fmt == FileFormat.DOCX else 3
    assert result['page_count'] == expected_pages
    metadata, displayed = assert_reader_matches_api(client, document_id, result, tmp_path)
    assert metadata['page_kind'] == ('slide' if fmt == FileFormat.PPTX else 'page')
    assert [p['page'] for p in result['extracted_text']] == list(range(1, expected_pages + 1))
    if fmt == FileFormat.PPTX:
        assert result['extracted_text'][2]['text'] == ''
    else:
        assert len(WordDocument(sources[fmt]).paragraphs) > expected_pages
    assert reader.extract_document_for_reader(document_id, session_factory, convert) == result
    assert len(calls) == 1  # No re-pagination or conversion on a revisit.
    assert client.get(metadata['reader_url']).content == displayed
    assert client.get(f'/documents/{document_id}/reader.pdf?version=wrong').status_code == 409


def test_conversion_failure_does_not_overwrite_an_existing_reader(client, session_factory, sources):
    source = sources[FileFormat.DOCX]
    document_id = save_document(session_factory, source, FileFormat.DOCX)
    old = reader.extract_document_for_reader(document_id, session_factory, fake_converter)
    old_url = client.get(f'/documents/{document_id}/pages').json()['reader_url']
    old_bytes = client.get(old_url).content
    doc = WordDocument(source)
    doc.add_paragraph('source changed')
    doc.save(source)
    def failing(*_):
        raise reader.ReaderConversionError('conversion_failed', 'Conversion failed explicitly')
    with pytest.raises(reader.ReaderConversionError, match='explicitly'):
        reader.extract_document_for_reader(document_id, session_factory, failing)
    assert client.get(old_url).content == old_bytes
    with session_factory() as db:
        assert db.get(Document, document_id).extracted_text == old['extracted_text']


def test_pptx_page_count_mismatch_is_explicit_and_saves_no_metadata(session_factory, sources):
    document_id = save_document(session_factory, sources[FileFormat.PPTX], FileFormat.PPTX)
    def wrong(source, fmt, output):
        return make_pdf(output / 'source.pdf', [b''])
    with pytest.raises(reader.ReaderConversionError) as error:
        reader.extract_document_for_reader(document_id, session_factory, wrong)
    assert error.value.code == 'slide_page_mismatch'
    with session_factory() as db:
        assert db.get(Document, document_id).page_count is None
        assert db.get(Document, document_id).extracted_text is None


def test_invalid_converted_pdf_and_commit_failure_clean_new_artifacts(session_factory, sources, monkeypatch):
    document_id = save_document(session_factory, sources[FileFormat.DOCX], FileFormat.DOCX)
    def invalid(source, fmt, output):
        path = output / 'invalid.pdf'
        path.write_bytes(b'not PDF')
        return path
    with pytest.raises(reader.ReaderConversionError) as error:
        reader.extract_document_for_reader(document_id, session_factory, invalid)
    assert error.value.code == 'invalid_rendered_pdf'
    def fail_commit(_):
        raise SQLAlchemyError('test DB failure')
    monkeypatch.setattr(session_factory.class_, 'commit', fail_commit)
    with pytest.raises(SQLAlchemyError):
        reader.extract_document_for_reader(document_id, session_factory, fake_converter)
    assert not list((sources[FileFormat.DOCX].parent / '.reader').rglob('*.pdf'))
    with session_factory() as db:
        assert db.get(Document, document_id).extracted_text is None


@pytest.mark.parametrize('suffix', ['pages', 'reader.pdf?version=' + '0' * 64])
def test_reader_authorization_precedes_file_access(client, session_factory, sources, suffix):
    document_id = save_document(session_factory, sources[FileFormat.DOCX], FileFormat.DOCX)
    other = f"Bearer {create_access_token({'sub': '124'})}"
    assert client.get(f'/documents/{document_id}/{suffix}', headers={'Authorization': other}).status_code == 403
    client.headers.pop('Authorization')
    assert client.get(f'/documents/{document_id}/{suffix}').status_code == 401


def test_prepare_endpoint_needs_no_external_program(client, session_factory, sources, monkeypatch):
    monkeypatch.setenv('PATH', '')
    document_id = save_document(session_factory, sources[FileFormat.DOCX], FileFormat.DOCX)
    response = client.post(f'/documents/{document_id}/reader')
    assert response.status_code == 200, response.text
    assert client.get(f'/documents/{document_id}/pages').json()['page_count'] == 2


def test_tampered_reader_pdf_is_not_served(client, session_factory, sources):
    document_id = save_document(session_factory, sources[FileFormat.DOCX], FileFormat.DOCX)
    reader.extract_document_for_reader(document_id, session_factory, fake_converter)
    url = client.get(f'/documents/{document_id}/pages').json()['reader_url']
    with session_factory() as db:
        path, _ = reader.reader_representation(db.get(Document, document_id))
    path.write_bytes(b'different file')
    assert client.get(url).status_code == 409
    reader.extract_document_for_reader(document_id, session_factory, fake_converter)
    assert client.get(url).status_code == 200


def test_corrupt_office_is_an_explicit_failure(tmp_path):
    source = tmp_path / 'broken.docx'
    source.write_bytes(b'not a document')
    with pytest.raises(reader.ReaderConversionError) as error:
        reader.convert_office(source, FileFormat.DOCX, tmp_path)
    assert error.value.code == 'conversion_failed'


@pytest.mark.parametrize('fmt', [FileFormat.DOCX, FileFormat.PPTX])
def test_native_office_pages_match_displayed_reader(client, session_factory, sources, fmt, tmp_path):
    document_id = save_document(session_factory, sources[fmt], fmt)
    response = client.post(f'/documents/{document_id}/reader')
    assert response.status_code == 200, response.text
    with session_factory() as db:
        doc = db.get(Document, document_id)
        result = {'page_count': doc.page_count, 'extracted_text': doc.extracted_text}
    metadata, _ = assert_reader_matches_api(client, document_id, result, tmp_path)
    assert metadata['reader_layout'] == 'text'
    assert metadata['file_format'] == fmt.value
    texts = [page['text'] for page in metadata['pages']]
    if fmt == FileFormat.DOCX:
        assert len(texts) == 2
        assert 'page one alpha' in texts[0] and 'page one beta' in texts[0]
        assert 'page two gamma' in texts[1] and 'page two delta' in texts[1]
        assert len(WordDocument(sources[fmt]).paragraphs) != len(texts)
    else:
        assert len(texts) == 3
        assert 'PPTX slide 1' in texts[0]
        assert 'PPTX slide 2' in texts[1]  # Hidden slide retains physical slide number 2.
        assert texts[2] == ''
        assert all('SPEAKER NOTES' not in text for text in texts)


def test_native_docx_preserves_vietnamese_tables_and_xml_characters(tmp_path):
    source = tmp_path / 'vietnamese.docx'
    document = WordDocument()
    document.add_paragraph('Thông tin hợp đồng: giá trị 1.234.567 VND, phạt 0.05%.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0,0).text = 'Thanh toán'
    table.cell(0,1).text = '30 ngày'
    table.cell(1,0).text = 'Ghi chú'
    table.cell(1,1).text = 'Kiểm tra & đối chiếu <bản gốc>'
    document.save(source)
    result = extract_pdf(reader.convert_office(source, FileFormat.DOCX, tmp_path))
    text = '\n'.join(p['text'] for p in result['extracted_text'])
    for expected in ['Thông tin hợp đồng', '1.234.567 VND', '0.05%', 'Thanh toán', '30 ngày', 'Kiểm tra & đối chiếu <bản gốc>']:
        assert expected in text
    assert all(page['text'][position['start']:position['end']] == position['text']
               for page in result['extracted_text'] for position in page['positions'])


def test_native_docx_flows_long_content_without_dropping_paragraphs(tmp_path):
    source = tmp_path / 'long.docx'
    document = WordDocument()
    for i in range(150):
        document.add_paragraph(f'SOURCE-{i:03d}: A full original paragraph with payment conditions and source details.')
    document.save(source)
    result = extract_pdf(reader.convert_office(source, FileFormat.DOCX, tmp_path))
    assert result['page_count'] > 1
    text = '\n'.join(p['text'] for p in result['extracted_text'])
    assert all(text.count(f'SOURCE-{i:03d}:') == 1 for i in range(150))


def test_native_pptx_preserves_grouped_text_and_dense_slide_on_one_page(tmp_path):
    source = tmp_path / 'grouped.pptx'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = 'Điều khoản & nghĩa vụ <gốc>'
    text_frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(3)).text_frame
    for i in range(80):
        text_frame.add_paragraph().text = f'ROW-{i:03d}: Original slide text with payment details.'
    deck.save(source)
    result = extract_pdf(reader.convert_office(source, FileFormat.PPTX, tmp_path))
    assert result['page_count'] == 1
    text = result['extracted_text'][0]['text']
    assert 'Điều khoản & nghĩa vụ <gốc>' in text
    assert all(text.count(f'ROW-{i:03d}:') == 1 for i in range(80))
